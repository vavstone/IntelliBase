"""Модель документов демо-корпуса и рендереры: DOCX (python-docx) и PDF (PyMuPDF).

Содержимое документов описывается декларативно — списком блоков (заголовок,
абзац, список, таблица, примечание). Один и тот же набор блоков умеет
отрисовать любой из двух рендереров, поэтому текст документа не зависит от
формата файла.

PDF собирается через `pymupdf.Story` (HTML/CSS → поток страниц A4). Кириллица
требует встроенного шрифта: используется DejaVu Sans (свободная лицензия,
полное покрытие кириллицы), путь ищется в системе, см. `find_font_dir`.
"""

from __future__ import annotations

import gc
import html
import os
from dataclasses import dataclass
from pathlib import Path

# --- блоки -----------------------------------------------------------------


@dataclass(frozen=True)
class Heading:
    """Заголовок раздела. level: 1 — раздел, 2 — подраздел, 3 — пункт."""

    level: int
    text: str


@dataclass(frozen=True)
class Para:
    """Абзац основного текста."""

    text: str


@dataclass(frozen=True)
class Bullets:
    """Маркированный список."""

    items: tuple[str, ...]


@dataclass(frozen=True)
class Numbers:
    """Нумерованный список."""

    items: tuple[str, ...]


@dataclass(frozen=True)
class Table:
    """Таблица с шапкой. `rows` — кортеж строк, в каждой столько же ячеек,
    сколько колонок в шапке."""

    header: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class Note:
    """Врезка «Примечание». В PDF — серый блок, в DOCX — курсивная строка."""

    text: str


Block = Heading | Para | Bullets | Numbers | Table | Note


@dataclass(frozen=True)
class Doc:
    """Документ демо-корпуса.

    category — slug категории (ПС), совпадает с каталогом в корпусе и с
    таксономией `kb_categories`; filename — имя файла с расширением
    (.docx или .pdf), расширение выбирает рендерер.
    """

    category: str
    filename: str
    title: str
    blocks: tuple[Block, ...]

    @property
    def suffix(self) -> str:
        return Path(self.filename).suffix.lower()


# --- шрифты -----------------------------------------------------------------

_FONT_SEARCH_DIRS: tuple[str, ...] = (
    "C:/Windows/Fonts",
    "/usr/share/fonts/truetype/dejavu",
    "/usr/share/fonts/dejavu",
    "/Library/Fonts",
)


def find_font_dir() -> Path:
    """Каталог с DejaVuSans.ttf. Переопределяется переменной DEMO_CORPUS_FONT_DIR."""

    env = os.environ.get("DEMO_CORPUS_FONT_DIR")
    candidates = ([env] if env else []) + list(_FONT_SEARCH_DIRS)
    for candidate in candidates:
        if candidate and (Path(candidate) / "DejaVuSans.ttf").exists():
            return Path(candidate)
    raise FileNotFoundError(
        "Не найден DejaVuSans.ttf (нужен для кириллицы в PDF). "
        "Укажите каталог со шрифтом в DEMO_CORPUS_FONT_DIR."
    )


# --- DOCX -------------------------------------------------------------------


def render_docx(doc: Doc, path: Path) -> None:
    """Записывает документ в .docx: Times New Roman 12 pt, таблицы с сеткой."""

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for block in doc.blocks:
        if isinstance(block, Heading):
            heading = document.add_heading(block.text, level=min(block.level, 4))
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif isinstance(block, Para):
            document.add_paragraph(block.text)
        elif isinstance(block, Bullets):
            for item in block.items:
                document.add_paragraph(item, style="List Bullet")
        elif isinstance(block, Numbers):
            for item in block.items:
                document.add_paragraph(item, style="List Number")
        elif isinstance(block, Table):
            table = document.add_table(rows=1, cols=len(block.header))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, block.header):
                cell.text = text
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in block.rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
        elif isinstance(block, Note):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"Примечание. {block.text}")
            run.italic = True

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


# --- PDF --------------------------------------------------------------------

_PDF_CSS = """
@font-face {font-family: "DejaVu"; src: url("DejaVuSans.ttf");}
@font-face {font-family: "DejaVu"; font-weight: bold; src: url("DejaVuSans-Bold.ttf");}
body {font-family: "DejaVu"; font-size: 10.5px; line-height: 1.45; color: #111;}
h1 {font-size: 16px; margin: 0 0 10px 0;}
h2 {font-size: 13px; margin: 16px 0 6px 0;}
h3 {font-size: 11.5px; margin: 12px 0 4px 0;}
p {margin: 0 0 8px 0; text-align: justify;}
ul, ol {margin: 0 0 10px 0; padding-left: 20px;}
li {margin-bottom: 3px;}
table {border-collapse: collapse; margin: 6px 0 12px 0; width: 100%;}
th, td {border: 0.6px solid #555; padding: 3px 6px; font-size: 9.5px; text-align: left;
        vertical-align: top;}
th {background: #ececec; font-weight: bold;}
.note {background: #f4f4f4; border-left: 3px solid #999; padding: 6px 9px;
       margin: 8px 0; font-size: 9.5px;}
"""


def _block_html(block: Block) -> str:
    if isinstance(block, Heading):
        level = min(max(block.level, 1), 3)
        return f"<h{level}>{html.escape(block.text)}</h{level}>"
    if isinstance(block, Para):
        return f"<p>{html.escape(block.text)}</p>"
    if isinstance(block, Bullets):
        items = "".join(f"<li>{html.escape(i)}</li>" for i in block.items)
        return f"<ul>{items}</ul>"
    if isinstance(block, Numbers):
        items = "".join(f"<li>{html.escape(i)}</li>" for i in block.items)
        return f"<ol>{items}</ol>"
    if isinstance(block, Table):
        head = "".join(f"<th>{html.escape(c)}</th>" for c in block.header)
        rows = "".join(
            "<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in row) + "</tr>"
            for row in block.rows
        )
        return f"<table><tr>{head}</tr>{rows}</table>"
    if isinstance(block, Note):
        return f'<div class="note">Примечание. {html.escape(block.text)}</div>'
    raise TypeError(f"Неизвестный блок: {type(block)!r}")


def render_pdf(doc: Doc, path: Path, font_dir: Path | None = None) -> None:
    """Записывает документ в .pdf потоком A4. Шрифт — DejaVu Sans из font_dir."""

    import pymupdf

    font_dir = font_dir or find_font_dir()
    body = "\n".join(_block_html(b) for b in doc.blocks)
    story = pymupdf.Story(
        html=f"<html><body>{body}</body></html>",
        user_css=_PDF_CSS,
        archive=pymupdf.Archive(str(font_dir)),
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    # Сборка идёт во временный файл: субсеттинг шрифтов выполняется отдельным
    # проходом, а на Windows нельзя перезаписать файл, пока он открыт.
    temp = path.with_suffix(path.suffix + ".raw")
    writer = pymupdf.DocumentWriter(str(temp))
    media_box = pymupdf.paper_rect("a4")
    where = media_box + (54, 54, -54, -54)  # поля 54 pt ≈ 19 мм
    more = True
    while more:
        device = writer.begin_page(media_box)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    # PyMuPDF на Windows отпускает дескриптор файла только при сборке мусора:
    # без этого временный файл не удалить.
    del writer
    gc.collect()

    _subset_fonts(temp, path)
    temp.unlink(missing_ok=True)


def _subset_fonts(source: Path, target: Path) -> None:
    """Оставляет только используемые глифы: без этого каждый PDF весит ~1,5 МБ
    (полные DejaVu Sans и DejaVu Sans Bold), после сжатия — десятки килобайт."""

    import pymupdf

    # Из памяти, а не с диска: PyMuPDF на Windows держит дескриптор открытого
    # файла до сборки мусора, из-за чего временный файл не удалить.
    document = pymupdf.open(stream=source.read_bytes(), filetype="pdf")
    try:
        document.subset_fonts()
        document.save(str(target), garbage=4, deflate=True, clean=True)
    finally:
        document.close()


def render(doc: Doc, out_dir: Path, font_dir: Path | None = None) -> Path:
    """Рендерит документ в out_dir/<category>/<filename> по расширению."""

    path = out_dir / doc.category / doc.filename
    if doc.suffix == ".docx":
        render_docx(doc, path)
    elif doc.suffix == ".pdf":
        render_pdf(doc, path, font_dir=font_dir)
    else:
        raise ValueError(f"Неподдерживаемое расширение: {doc.filename}")
    return path
